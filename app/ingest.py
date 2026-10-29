import io
import zipfile
from pathlib import Path

from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.errors import AppError

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}
SPLITTER = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=180)


def extract_chunks(filename: str, raw: bytes, max_characters: int) -> list[dict]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED:
        raise AppError("Choose a PDF, DOCX, TXT, or Markdown file.", 415)
    if not raw:
        raise AppError("This file is empty.")
    pages: list[tuple[int, str]] = []
    try:
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise AppError(
                    "Password-protected PDFs are not supported. Upload an unlocked copy."
                )
            if len(reader.pages) > 500:
                raise AppError("Use a PDF with at most 500 pages.")
            total = 0
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                total += len(text)
                if total > max_characters:
                    raise AppError("The extracted document exceeds the 1 million character limit.")
                pages.append((number, text))
        elif extension == ".docx":
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise AppError("The expanded DOCX exceeds the 50 MB limit.")
            doc = Document(io.BytesIO(raw))
            lines = [p.text for p in doc.paragraphs]
            lines.extend(
                " | ".join(cell.text for cell in row.cells)
                for table in doc.tables
                for row in table.rows
            )
            pages = [(1, "\n".join(lines))]
        else:
            pages = [(1, raw.decode("utf-8-sig"))]
    except AppError:
        raise
    except UnicodeDecodeError as exc:
        raise AppError("Save text files using UTF-8 encoding and upload again.") from exc
    except Exception as exc:
        raise AppError(
            "Unable to read this file. Check that it is a valid, unlocked document."
        ) from exc
    if sum(len(text) for _, text in pages) > max_characters:
        raise AppError("The extracted document exceeds the 1 million character limit.")
    chunks = []
    for page, text in pages:
        for part in SPLITTER.split_text(text.replace("\x00", "")):
            if part.strip():
                chunks.append({"page": page if extension == ".pdf" else None, "text": part})
    if not chunks:
        raise AppError("No readable text found. Scanned PDFs need OCR before uploading.")
    return chunks
