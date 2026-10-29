import io
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import Settings
from app.errors import AppError
from app.ingest import extract_chunks
from app.main import create_app
from tests.conftest import FakeEmbeddings, FakeGrok


def upload(client, name="policy.txt", text=b"Refund requests are accepted within 30 days."):
    return client.post("/api/documents", files={"file": (name, text, "text/plain")})


def test_upload_ask_delete_and_citations(client):
    response = upload(client)
    assert response.status_code == 201
    identifier = response.json()["document"]["id"]
    assert client.get("/api/documents").json()["documents"][0]["chunks"] == 1
    answer = client.post("/api/ask", json={"question": "What is the refund policy?"})
    assert answer.status_code == 200
    data = answer.json()
    assert "30 days" in data["answer"]
    assert data["sources"][0]["cited"]
    assert data["sources"][0]["document_id"] == identifier
    assert data["sources"][0]["page"] is None
    assert client.delete(f"/api/documents/{identifier}").status_code == 204
    assert client.get("/api/documents").json()["documents"] == []
    assert client.post("/api/ask", json={"question": "Refund?"}).status_code == 409
    assert client.delete(f"/api/documents/{identifier}").status_code == 404


def test_duplicate_and_path_filename(client):
    first = upload(client, "../../policy.txt").json()
    second = upload(client, "another.txt").json()
    assert first["document"]["name"] == "policy.txt"
    assert second["duplicate"]
    assert first["document"]["id"] == second["document"]["id"]
    assert len(client.get("/api/documents").json()["documents"]) == 1


def test_selected_document_filter(client):
    refund = upload(client).json()["document"]["id"]
    shipping = upload(client, "delivery.txt", b"Shipping takes 5 business days.").json()[
        "document"
    ]["id"]
    answer = client.post(
        "/api/ask", json={"question": "Refund?", "document_ids": [shipping]}
    ).json()
    assert {s["document_id"] for s in answer["sources"]} == {shipping}
    best = client.post("/api/ask", json={"question": "Refund?", "top_k": 1}).json()
    assert best["sources"][0]["document_id"] == refund
    client.delete(f"/api/documents/{shipping}")
    assert (
        client.post("/api/ask", json={"question": "?", "document_ids": [shipping]}).status_code
        == 404
    )


@pytest.mark.parametrize(
    "name,raw,code",
    [
        ("run.exe", b"hello", 415),
        ("empty.txt", b"", 400),
        ("invalid.pdf", b"not a pdf", 400),
        ("invalid.docx", b"not a zip", 400),
        ("bad.txt", b"\xff\xfe", 400),
        ("blank.txt", b"   \n  ", 400),
    ],
)
def test_invalid_documents(client, name, raw, code):
    assert upload(client, name, raw).status_code == code


@pytest.mark.parametrize(
    "body",
    [
        {"question": " "},
        {"question": "a" * 2001},
        {"question": "?", "top_k": 9},
        {"question": "?", "document_ids": ["bad-id"]},
    ],
)
def test_question_validation(client, body):
    assert client.post("/api/ask", json=body).status_code == 422


def test_missing_key(tmp_path):
    settings = Settings(
        _env_file=None, data_dir=tmp_path / "no-key", llm_provider="grok", xai_api_key=""
    )
    with TestClient(create_app(settings, FakeEmbeddings())) as client:
        assert client.get("/api/health").json()["grok_configured"] is False
        upload(client)
        response = client.post("/api/ask", json={"question": "Refund?"})
        assert response.status_code == 503
        assert "XAI_API_KEY" in response.json()["detail"]


def test_persistence_after_restart(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path / "saved")
    with TestClient(create_app(settings, FakeEmbeddings(), FakeGrok())) as client:
        identifier = upload(client).json()["document"]["id"]
    with TestClient(create_app(settings, FakeEmbeddings(), FakeGrok())) as client:
        assert client.get("/api/documents").json()["documents"][0]["id"] == identifier
        assert "30 days" in client.post("/api/ask", json={"question": "Refund?"}).json()["answer"]


def test_docx_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Our refund policy.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Refund window"
    table.cell(0, 1).text = "30 days"
    buffer = io.BytesIO()
    doc.save(buffer)
    chunks = extract_chunks("test.docx", buffer.getvalue(), 1_000_000)
    assert "30 days" in chunks[0]["text"]
    assert "Our refund" in chunks[0]["text"]


def test_oversized_docx_expansion():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"a" * (51 * 1024 * 1024))
    with pytest.raises(AppError, match="expanded DOCX"):
        extract_chunks("large.docx", buffer.getvalue(), 1_000_000)


def test_text_limit_and_chunk_overlap():
    with pytest.raises(AppError, match="character limit"):
        extract_chunks("big.txt", b"a" * 1001, 1000)
    chunks = extract_chunks("long.txt", ("refund policy applies. " * 200).encode(), 1_000_000)
    assert len(chunks) > 1
    assert all(len(c["text"]) <= 1000 for c in chunks)


def test_limits_and_malformed_content_length(tmp_path):
    settings = Settings(
        _env_file=None, data_dir=tmp_path / "limits", max_upload_bytes=100, max_documents=1
    )
    with TestClient(create_app(settings, FakeEmbeddings(), FakeGrok())) as client:
        assert upload(client, text=b"a" * 101).status_code == 413
        assert upload(client).status_code == 201
        assert upload(client, "second.txt", b"different text").status_code == 409
        assert (
            client.post(
                "/api/documents", content=b"", headers={"content-length": "99999999"}
            ).status_code
            == 413
        )
        assert (
            client.post(
                "/api/documents", content=b"", headers={"content-length": "invalid"}
            ).status_code
            == 413
        )


def test_static_and_security_headers(client):
    response = client.get("/")
    assert response.status_code == 200
    assert "Folio" in response.text
    assert "script-src 'self'" in response.headers["Content-Security-Policy"]
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/openapi.json").status_code == 200


def test_invalid_citations(tmp_path):
    class InvalidGrok:
        def answer(self, question, sources):
            return "Made up citation [S999]"

    settings = Settings(_env_file=None, data_dir=tmp_path / "citations")
    with TestClient(create_app(settings, FakeEmbeddings(), InvalidGrok())) as client:
        upload(client)
        assert client.post("/api/ask", json={"question": "Refund?"}).status_code == 502
